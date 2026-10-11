# -*- coding: utf-8 -*-
"""按文档顺序读取 docx 的正文段落与表格单元格，供 style_lint.py 与 polish_diff.py 共用。

    read_docx(path) -> [block, ...]

block 为 dict：
  kind   'para'：正文层段落（含标题、图题表题、表注、参考文献与空段）
         'cell'：表格单元格（合并单元格只取一次）
  n      para：正文层段落序号，从 1 起，含空段，与 python-docx 的 Document.paragraphs 同序；
         cell：表格之前最后一个正文层段落的序号，用来把表格锚回正文位置
  loc    定位文字：'P12'，或 '表3 r2c3'（表号取表格前最近的「表N」题注，找不到按出现次序编号）
  text   run 文字按序拼接（制表符记 \\t，换行记 \\n；单元格内多段以 \\n 连接）
  cites  [(引用文字, 是否上标), ...]，引用文字形如 [12]、[6-7]、[5,19,22,34-35]
  sups   段内全部上标 run 的文字，按出现次序

只依赖 python-docx；文本框、页眉页脚与脚注不读取。上标只认直接格式（w:vertAlign），
不追溯字符样式。
"""
import re

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph

CITE_TXT = re.compile(r'\[\d+(?:\s*[-\u2013\u2014,\uff0c]\s*\d+)*\]')
TABLE_CAPTION = re.compile(r'^\s*表\s*(\d+)')


def _runs(par):
    for item in par.iter_inner_content():
        if isinstance(item, Hyperlink):
            for r in item.runs:
                yield r
        else:
            yield item


def _is_sup(run):
    va = run._r.find('./' + qn('w:rPr') + '/' + qn('w:vertAlign'))
    return va is not None and va.get(qn('w:val')) == 'superscript'


def _read_par(par):
    """返回 (text, cites, sups)；cites 的上标判定按字符区间落在哪个 run 上。"""
    pieces, spans, sups = [], [], []
    pos = 0
    for r in _runs(par):
        t = r.text
        if not t:
            continue
        sup = _is_sup(r)
        pieces.append(t)
        spans.append((pos, pos + len(t), sup))
        if sup:
            sups.append(t)
        pos += len(t)
    text = ''.join(pieces)
    cites = []
    for m in CITE_TXT.finditer(text):
        sup = any(s <= m.start() < e and flag for s, e, flag in spans)
        cites.append((m.group(0), sup))
    return text, cites, sups


def read_docx(path):
    doc = Document(path)
    blocks = []
    n = 0
    n_table = 0
    last_caption = None
    for child in doc.element.body.iterchildren():
        if child.tag == qn('w:p'):
            n += 1
            text, cites, sups = _read_par(Paragraph(child, doc))
            m = TABLE_CAPTION.match(text)
            if m:
                last_caption = m.group(1)
            blocks.append(dict(kind='para', n=n, loc='P%d' % n, text=text,
                               cites=cites, sups=sups))
        elif child.tag == qn('w:tbl'):
            n_table += 1
            label = last_caption or str(n_table)
            last_caption = None
            # lxml 代理对象被回收后 id 会复用，必须持有引用才能按 id 去重合并单元格
            held, seen = [], set()
            for ri, row in enumerate(Table(child, doc).rows, 1):
                for ci, cell in enumerate(row.cells, 1):
                    if id(cell._tc) in seen:
                        continue
                    held.append(cell._tc)
                    seen.add(id(cell._tc))
                    texts, cites, sups = [], [], []
                    for par in cell.paragraphs:
                        t, c, s = _read_par(par)
                        texts.append(t)
                        cites.extend(c)
                        sups.extend(s)
                    blocks.append(dict(kind='cell', n=n,
                                       loc='表%s r%dc%d' % (label, ri, ci),
                                       text='\n'.join(texts), cites=cites,
                                       sups=sups))
    return blocks
